import openpyxl, json, pathlib
here = pathlib.Path(__file__).parent
D = pathlib.Path(r'C:\Users\0029\Desktop\確定仕様書読み解き')
ws = openpyxl.load_workbook(D / 'HRB確定仕様書読み解き.xlsx', data_only=True)['OP一覧']
ops = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None: continue
    grp = str(r[1]).replace('HEB', 'HRB')
    g = 'HRB小型・中型' if ('HRB中型' in grp or 'HRB小型' in grp) else grp
    ops.append({'no': r[0], 'grp': g, 'name': str(r[2]).replace('\n', ''), 'cond': str(r[3] or ''), 'names': [str(x) for x in r[4:7] if x]})
(here / 'ocrlib' / 'oplist.json').write_text(json.dumps(ops, ensure_ascii=False), encoding='utf-8')
ws = openpyxl.load_workbook(D / '確定仕様書読み解き HRB 2026年12月（最終テスト）.xlsx', data_only=True).active
names = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
rows = []
for r in range(4, ws.max_row + 1):
    v = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    rows.append({'ym': v[0], 'name': v[1], 'ser': str(v[2]), 'shimuke': v[3], 'country': v[4], 'cust': v[5],
                 'ops': [str(names[i]) for i in range(6, len(v) - 1) if v[i] == 1], 'biko': v[-1] or ''})
(here / 'ocrlib' / 'answer.json').write_text(json.dumps(rows, ensure_ascii=False), encoding='utf-8')
print(len(ops), len(rows), rows[0])
