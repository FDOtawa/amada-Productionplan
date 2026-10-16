"""12月の確定仕様書（HRB・KD・SRB）を目で読んだ結果から、OP一覧のきまりで帳票（試し）を作る"""
import re, unicodedata, pathlib, openpyxl
from openpyxl.styles import Alignment, Border, Side, PatternFill, Font
exec(open(pathlib.Path(__file__).parent / 'dec_read.py', encoding='utf-8').read())
LARGE = [dict(f='組付(中型)', p=42, ym='2026年12月', ser='35065013', model='HRB3506', country='日本', dest='国内',
              cust='(株)ソーゴ', contract='(株)ソーゴ', tok='BD1-4872',
              on=['AS-01', '後方LED', 'ｸﾗﾝﾌﾟOP'],
              note='特殊仕様①〜⑧（ビットレス・アンカー対応ベースプレート・特殊ワークセンササツキアテ等）／改訂A：突き当て先端厚特殊30mm を確定欄で選択しているが不要／製造指示：80周年メモリアルプレート')]
SRB = [dict(f='SRB', p=1, ym='2026年12月', ser='10301114', model='SRB1003', country='インド', dest='オーバーシーズ',
            cust='Amada (India) Pvt. Ltd. - Technical Center', contract='Amada (India) Pvt. Ltd. - Technical Center', tok='',
            on=['LED照明', 'ﾃﾊﾟBOX', 'ﾊﾞｯｸｹﾞｰｼﾞ4軸'], note='', base='国内')]
OPS = [r for r in openpyxl.load_workbook(r'C:\Users\0029\Desktop\HRB確定仕様書読み解き.xlsx', data_only=True)['OP一覧'].iter_rows(min_row=2, values_only=True) if r[0] is not None]

KANA = str.maketrans({})
def hankaku(s):
    out = ''
    for ch in s:
        n = unicodedata.normalize('NFKC', ch) if not ('ァ' <= ch <= 'ヶ' or ch in 'ー・') else ch
        out += n
    # 全角カタカナ → 半角カタカナ（濁点・半濁点も分ける）
    table = {}
    full = 'アイウエオカキクケコサシスセソタチツテトナニヌネノハヒフヘホマミムメモヤユヨラリルレロワヲンァィゥェォャュョッー・'
    half = 'ｱｲｳｴｵｶｷｸｹｺｻｼｽｾｿﾀﾁﾂﾃﾄﾅﾆﾇﾈﾉﾊﾋﾌﾍﾎﾏﾐﾑﾒﾓﾔﾕﾖﾗﾘﾙﾚﾛﾜｦﾝｧｨｩｪｫｬｭｮｯｰ･'
    for a, b in zip(full, half): table[a] = b
    for a, b in zip('ガギグゲゴザジズゼゾダヂヅデドバビブベボ', 'ｶｷｸｹｺｻｼｽｾｿﾀﾁﾂﾃﾄﾊﾋﾌﾍﾎ'): table[a] = b + 'ﾞ'
    for a, b in zip('パピプペポ', 'ﾊﾋﾌﾍﾎ'): table[a] = b + 'ﾟ'
    table['ヴ'] = 'ｳﾞ'
    return ''.join(table.get(c, c) for c in out)

def tons(model):
    d = re.sub(r'[^0-9]', '', re.sub(r'^(HRBC|HRB|SRB)', '', model))[:4]
    return int(d[:2]) * (10 if int(d[:2]) >= 10 and d[2] == '0' and int(d[:3]) in (100, 103, 130, 220, 350, 500, 600) else 1) if False else \
        {'1003': 100, '1303': 130, '2204': 220, '5020': 50, '8025': 80, '3503': 350, '3504': 350, '3506': 350, '6007': 600}.get(d)

rows = []
for x in LARGE:
    if x.get('skip'): continue
    t = tons(x['model'])
    is_srb = x['model'].startswith('SRB')
    kind = ('SRB中型' if t == 220 else 'SRB小型') if is_srb else ('HRB大型' if t >= 350 else 'HRB中型' if t == 220 else 'HRB小型')
    EU = {'欧州', 'ドイツ', 'フランス', 'イタリア', 'スペイン', 'イギリス', 'オランダ', 'ベルギー', 'スイス', 'オーストリア', 'ポーランド', 'チェコ', 'スウェーデン', 'デンマーク', 'フィンランド', 'ノルウェー', 'ポルトガル', 'ハンガリー', 'トルコ'}
    shimuke = 'DMS' if x['country'] == '日本' else ('CE' if is_srb and x['country'] in EU else 'EXP')  # 国名を最優先（2026-10-16 ユーザー回答）
    aai = x['cust'] == 'AAI'
    name = f"{x['model']} {shimuke}" + (' AAI' if aai else '')
    on = set(x['on']); atc = 'ATC' in x['model']
    small, mid = kind == 'HRB小型', kind == 'HRB中型'
    hrb_sm = small or mid
    f = {}
    if hrb_sm:
        f[1] = atc
        f[2] = small and 'ｵｰﾄｸﾗｳﾆﾝｸﾞ' in on
        f[3] = mid and 'ｵｰﾄｸﾗｳﾆﾝｸﾞ' in on
        f[4] = 'Lｼﾌﾄ' in on
        f[5] = 'BI-S1' in on; f[6] = 'BI-S2' in on
        f[7] = 'AS-01' in on; f[8] = 'AS-01' not in on  # W24 は大型にも付く（OP一覧 2026-10-16）
        ag = 'AGRIP-A' in on
        f[9] = ag and small and t != 130 and shimuke == 'DMS'; f[10] = ag and t == 130 and shimuke == 'DMS'; f[11] = ag and mid and shimuke == 'DMS'
        f[12] = ag and small and t != 130 and shimuke == 'EXP'; f[13] = ag and t == 130 and shimuke == 'EXP'; f[14] = ag and mid and shimuke == 'EXP'
        f[15] = '自動ｽﾗｲﾄﾞﾌｯﾄ' in on; f[16] = '手動ｽﾗｲﾄﾞﾌｯﾄ' in on; f[17] = 'ﾛﾎﾞｯﾄI/F' in on; f[18] = '二人ﾌｯﾄ' in on; f[19] = '20P+AF' in on
        f[20] = atc
        f[21] = '追従装置SF1224' in on and (small or (mid and shimuke == 'DMS')); f[22] = '追従装置SF1224' in on and mid and shimuke == 'EXP'
        f[23] = '追従装置SF1584' in on; f[24] = 'VFﾕﾆｯﾄ' in on
        f[25] = '制御ASSY(IO)' in on and shimuke == 'DMS'; f[26] = '制御ASSY(IO)' in on and shimuke == 'EXP'
        f[42] = aai
    if kind == 'HRB大型':
        f[8] = 'AS-01' not in on  # W24 は大型にも（OP一覧 2026-10-16）
        f[29] = 'ﾍﾋﾞｰﾊﾞｯｸｹﾞｰｼﾞY4軸' in on; f[30] = 'BI-L1' in on; f[31] = 'BI-L2' in on; f[32] = '後方LED' in on; f[33] = 'AS-01' in on
        f[34] = 'ｸﾗﾝﾌﾟOP' in on; f[35] = 'ﾛﾎﾞｯﾄI/F' in on; f[36] = 'ﾜｰｸｾﾝｻ' in on
        f[37] = 'QF1' not in on; f[38] = 'QF1' in on; f[39] = '制御ASSY(IO)' in on; f[40] = '二人ﾌｯﾄ' in on; f[41] = atc
    f[27] = '指定色' in on; f[28] = bool(x['tok'])
    if is_srb:
        exp, ce = shimuke == 'EXP', shimuke == 'CE'
        f[43] = exp and 'AS-01' in on; f[44] = ce and 'AS-01' in on; f[45] = exp and 'AS-01' not in on; f[48] = ce and 'AS-01' not in on
        f[49] = exp and 'LED照明' in on; f[50] = ce and 'LED照明' in on; f[51] = 'ﾊﾞｯｸｹﾞｰｼﾞ4軸' in on; f[52] = 'BI-S1' in on; f[53] = ce and 'ﾃﾊﾟBOX' in on
    biko = []
    if x.get('mismatch'): biko.append('ｱﾝﾏｯﾁ有')
    if '追従装置SF1584' in on: biko.append('追従装置SF1584 選択あり')
    biko = '・'.join(biko)
    rows.append(dict(x=x, name=name, shimuke=shimuke, kind=kind, f=f, biko=biko,
                     cust=hankaku(x['cust'] or x['contract'])))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = '確定仕様書読み解き'
base = ['ｱﾏﾀﾞ生産月', '機種名', 'ｼﾘｱﾙ', '仕向', '国名', '客先']
thin = Side(style='thin', color='A6A6A6'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
head = PatternFill('solid', fgColor='DDEBF7'); grp = {'HRB小型・中型': 'FFF2CC', 'HRB大型': 'E2EFDA', 'SRB': 'FCE4D6', '全品': 'EDEDED'}
for i, h in enumerate(base, 1):
    ws.cell(1, i, h); ws.merge_cells(start_row=1, start_column=i, end_row=3, end_column=i)
c0 = len(base) + 1
ws.cell(1, c0, 'OP'); ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + len(OPS) - 1)
gname = lambda b: 'HRB小型・中型' if 'HEB' in b or 'HRB中型' in b else b
col = c0; prev = None; start = c0
for op in OPS:
    g = gname(op[1])
    if g != prev:
        if prev: ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=col - 1)
        ws.cell(2, col, g); prev, start = g, col
    ws.cell(3, col, str(op[2]).replace('\n', ''))
    col += 1
ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=col - 1)
ws.cell(1, col, '備考欄'); ws.merge_cells(start_row=1, start_column=col, end_row=3, end_column=col)
last = col
for r, row in enumerate(rows, 4):
    x = row['x']
    vals = [x['ym'], row['name'], x['ser'], row['shimuke'], {'アラブ首長国連邦': 'UAE'}.get(x['country'], x['country']), row['cust']]
    for i, v in enumerate(vals, 1): ws.cell(r, i, v)
    for k, op in enumerate(OPS):
        if row['f'].get(op[0]): ws.cell(r, c0 + k, 1)
    ws.cell(r, last, row['biko'])
for r in range(1, len(rows) + 4):
    for c in range(1, last + 1):
        cell = ws.cell(r, c); cell.border = bd
        if r <= 3:
            cell.fill = head; cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        elif c >= c0 and c < last: cell.alignment = Alignment(horizontal='center')
for c in range(c0, last):
    g = ws.cell(2, c).value
for k, op in enumerate(OPS):
    ws.cell(3, c0 + k).alignment = Alignment(horizontal='center', vertical='center', text_rotation=255, wrap_text=True)
    ws.cell(3, c0 + k).fill = PatternFill('solid', fgColor=grp.get(gname(op[1]), 'DDEBF7'))
    ws.column_dimensions[openpyxl.utils.get_column_letter(c0 + k)].width = 4.2
for c in range(c0, last):
    v = ws.cell(2, c).value
    if v: ws.cell(2, c).fill = PatternFill('solid', fgColor=grp.get(v, 'DDEBF7'))
ws.row_dimensions[3].height = 150
for i, w in enumerate([12, 22, 11, 6, 15, 40], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.column_dimensions[openpyxl.utils.get_column_letter(last)].width = 16
ws.freeze_panes = ws.cell(4, c0)
ws.page_setup.paperSize = 9; ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '1:3'
ws.oddFooter.center.text = '&P / &N'
out = pathlib.Path(r'C:\Users\0029\Desktop\確定仕様書読み解き HRB大型 2026年12月（試し）.xlsx')
n = 1
while out.exists(): out = out.with_name(f'確定仕様書読み解き HRB大型 2026年12月（試し） ({n}).xlsx'); n += 1
wb.save(out); print(out, len(rows))
for row in rows:
    print(row['x']['f'], row['x']['p'], row['name'], row['x']['ser'], row['kind'], row['shimuke'], ' '.join(str(k) for k, v in sorted(row['f'].items()) if v), row['biko'])
